#!/usr/bin/env python3
"""
Generate Standard Word Document (.docx) adhering to Egyptian Pharaoh Academic Design System
Author: Howard Liao Ph.D. (廖倫豪 博士)
"""

import os
import docx
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

OUTPUT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_FILE = os.path.join(OUTPUT_DIR, "Pharaoh_Academic_Report_HowardLiao.docx")

# Colors
HEX_PRIMARY = "D96B27"       # 暖赭石橙
HEX_SECONDARY = "2B5F6B"     # 綠松石青
HEX_TERTIARY = "E5A93C"      # 法老琥珀金
HEX_NEUTRAL = "1F2421"       # 曜石深灰
HEX_MUTED = "6B726C"         # 沉香灰
HEX_SURFACE_WARM = "FAF7F2"  # 羊皮紙暖白
HEX_BORDER = "E8DFD8"        # 暖沙灰
HEX_ACCENT_SOFT = "FEF7E9"   # 琥珀柔光底

RGB_PRIMARY = RGBColor(217, 107, 39)
RGB_SECONDARY = RGBColor(43, 95, 107)
RGB_TERTIARY = RGBColor(229, 169, 60)
RGB_NEUTRAL = RGBColor(31, 36, 33)
RGB_MUTED = RGBColor(107, 114, 108)

def set_cell_background(cell, hex_color):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_report():
    doc = docx.Document()

    # Set standard margins (1 inch / 2.54 cm)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

        # Header
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("埃及法老溫暖學術風格 | 規範與實務指南")
        hrun.font.name = "Outfit"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGB_MUTED

        # Footer
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        frun_l = fp.add_run("報告人 Howard Liao Ph.D.(廖倫豪 博士)  |  學術與專業研究")
        frun_l.font.name = "Inter"
        frun_l.font.size = Pt(8.5)
        frun_l.font.color.rgb = RGB_MUTED

    # ==========================================
    # 封面 (Cover Page)
    # ==========================================
    p_pre = doc.add_paragraph()
    p_pre.paragraph_format.space_before = Pt(72)
    p_pre.paragraph_format.space_after = Pt(8)
    r_tag = p_pre.add_run("PROFESSIONAL ARCHITECTURE & RESEARCH REPORT")
    r_tag.font.name = "Outfit"
    r_tag.font.size = Pt(11)
    r_tag.font.bold = True
    r_tag.font.color.rgb = RGB_SECONDARY

    # Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_after = Pt(12)
    r_title = p_title.add_run("埃及法老溫暖學術風格：\n跨平台統一設計系統實踐白皮書")
    r_title.font.name = "Outfit"
    r_title.font.size = Pt(28)
    r_title.font.bold = True
    r_title.font.color.rgb = RGB_PRIMARY

    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(36)
    r_sub = p_sub.add_run("統一 PPT、Word、Google Docs 及 Web 前端之視覺語彙與自動化生成規範")
    r_sub.font.name = "Inter"
    r_sub.font.size = Pt(14)
    r_sub.font.color.rgb = RGB_NEUTRAL

    # Decorative line (Table with 1 row, 2 cells: orange & gold)
    line_tbl = doc.add_table(rows=1, cols=2)
    line_tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_cell_background(line_tbl.rows[0].cells[0], HEX_PRIMARY)
    set_cell_background(line_tbl.rows[0].cells[1], HEX_TERTIARY)
    line_tbl.rows[0].cells[0].width = Inches(2.0)
    line_tbl.rows[0].cells[1].width = Inches(4.5)
    # Make it very thin
    trPr = line_tbl.rows[0]._element.get_or_add_trPr()
    trHeight = parse_xml(f'<w:trHeight {nsdecls("w")} w:val="50" w:hRule="exact"/>')
    trPr.append(trHeight)

    # Attribution Callout Card (Bottom of Cover)
    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(120)

    auth_tbl = doc.add_table(rows=1, cols=1)
    auth_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_auth = auth_tbl.rows[0].cells[0]
    c_auth.width = Inches(6.5)
    set_cell_background(c_auth, HEX_ACCENT_SOFT)
    set_cell_margins(c_auth, top=200, bottom=200, left=240, right=240)

    # Border for auth card
    tcPr = c_auth._element.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{HEX_PRIMARY}"/>'
        f'<w:top w:val="single" w:sz="8" w:space="0" w:color="{HEX_BORDER}"/>'
        f'<w:right w:val="single" w:sz="8" w:space="0" w:color="{HEX_BORDER}"/>'
        f'<w:bottom w:val="single" w:sz="8" w:space="0" w:color="{HEX_BORDER}"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)

    p_a1 = c_auth.paragraphs[0]
    p_a1.paragraph_format.space_after = Pt(4)
    ra1 = p_a1.add_run("報告人 Howard Liao Ph.D.(廖倫豪 博士)")
    ra1.font.name = "Outfit"
    ra1.font.size = Pt(14)
    ra1.font.bold = True
    ra1.font.color.rgb = RGB_PRIMARY

    p_a2 = c_auth.add_paragraph()
    p_a2.paragraph_format.space_after = Pt(0)
    ra2 = p_a2.add_run("資深架構師 / 資訊安全長 (CISO) / 資訊工程博士  |  發布版本：v1.0 (2026)")
    ra2.font.name = "Inter"
    ra2.font.size = Pt(9.5)
    ra2.font.color.rgb = RGB_SECONDARY

    # Page Break to Content
    doc.add_page_break()

    # ==========================================
    # 內文 (Content Page)
    # ==========================================
    # Heading 1
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(8)
    rh1 = h1.add_run("第一章 設計哲學與核心原則")
    rh1.font.name = "Outfit"
    rh1.font.size = Pt(18)
    rh1.font.bold = True
    rh1.font.color.rgb = RGB_PRIMARY

    # Body Paragraph
    p_body1 = doc.add_paragraph()
    p_body1.paragraph_format.space_after = Pt(10)
    p_body1.paragraph_format.line_spacing = 1.35
    rb1 = p_body1.add_run(
        "「埃及法老溫暖學術風格」融合了古文明典籍的永恆厚重與現代數位工程的極簡幾何。傳統文件生成常面臨全黑文字過度刺眼、純白畫布缺乏溫度、或跨簡報與文件時色彩破碎等痛點。本設計系統以羊皮紙暖白（#FAF7F2）為底色載體，搭配曜石深灰（#1F2421）作為高對比正文字體，徹底達到護眼與學術莊重的雙重平衡。"
    )
    rb1.font.name = "Inter"
    rb1.font.size = Pt(10.5)
    rb1.font.color.rgb = RGB_NEUTRAL

    # Heading 2
    h2 = doc.add_paragraph()
    h2.paragraph_format.space_before = Pt(14)
    h2.paragraph_format.space_after = Pt(6)
    rh2 = h2.add_run("1.1 核心調色盤與語意對應")
    rh2.font.name = "Outfit"
    rh2.font.size = Pt(14)
    rh2.font.bold = True
    rh2.font.color.rgb = RGB_SECONDARY

    # Table with styling
    tbl = doc.add_table(rows=5, cols=4)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

    headers = ["語意名稱", "十六進位碼 (Hex)", "RGB 色彩值", "標準應用情境"]
    for i, h_text in enumerate(headers):
        c = tbl.rows[0].cells[i]
        set_cell_background(c, HEX_PRIMARY)
        set_cell_margins(c, top=140, bottom=140, left=160, right=160)
        p = c.paragraphs[0]
        run = p.add_run(h_text)
        run.font.name = "Outfit"
        run.font.size = Pt(10)
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)

    data = [
        ("Primary 暖赭石橙", "#D96B27", "217, 107, 39", "主標題、核心關鍵字、封面首要引導"),
        ("Secondary 綠松石青", "#2B5F6B", "43, 95, 107", "副標題、次級章節、專業標籤、架構分區"),
        ("Tertiary 法老琥珀金", "#E5A93C", "229, 169, 60", "KPI 高亮、決策重點卡片邊框、徽章亮色"),
        ("Neutral 曜石深灰", "#1F2421", "31, 36, 33", "主要內文正文、高可讀性段落文字")
    ]

    for row_idx, row_data in enumerate(data, start=1):
        bg_color = HEX_SURFACE_WARM if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(row_data):
            c = tbl.rows[row_idx].cells[col_idx]
            set_cell_background(c, bg_color)
            set_cell_margins(c, top=120, bottom=120, left=140, right=140)
            p = c.paragraphs[0]
            run = p.add_run(text)
            run.font.name = "Inter"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGB_NEUTRAL
            if col_idx == 0:
                run.font.bold = True

    # Callout Highlight Box
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(14)
    p_sp.paragraph_format.space_after = Pt(4)

    callout_tbl = doc.add_table(rows=1, cols=1)
    callout_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_box = callout_tbl.rows[0].cells[0]
    c_box.width = Inches(6.5)
    set_cell_background(c_box, HEX_ACCENT_SOFT)
    set_cell_margins(c_box, top=160, bottom=160, left=200, right=200)

    # Amber Left Border
    tcPr = c_box._element.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="28" w:space="0" w:color="{HEX_TERTIARY}"/>'
        f'<w:top w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)

    cp = c_box.paragraphs[0]
    c_run1 = cp.add_run("💡 法定落款與輸出規範提示：\n")
    c_run1.font.name = "Outfit"
    c_run1.font.size = Pt(10.5)
    c_run1.font.bold = True
    c_run1.font.color.rgb = RGB_PRIMARY

    c_run2 = cp.add_run(
        "所有透過 HermesAgent 產出之研究白皮書、架構評估案及策略提案，在封面及頁尾必須嚴格保留『報告人 Howard Liao Ph.D.(廖倫豪 博士)』標準字樣，以確保智慧產權標註完整與高階學術權威信譽。"
    )
    c_run2.font.name = "Inter"
    c_run2.font.size = Pt(9.5)
    c_run2.font.color.rgb = RGB_NEUTRAL

    doc.save(OUTPUT_FILE)
    print(f"DOCX generated successfully at: {OUTPUT_FILE}")

if __name__ == "__main__":
    create_report()
