#!/usr/bin/env python3
"""
Generate Cloud Architecture Strategic Deck
adhering to Egyptian Pharaoh Academic Design System
Author: Howard Liao Ph.D. (廖倫豪 博士)
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

COLOR_PRIMARY = RGBColor(217, 107, 39)       # 暖赭石橙 #D96B27
COLOR_SECONDARY = RGBColor(43, 95, 107)     # 綠松石青 #2B5F6B
COLOR_TERTIARY = RGBColor(229, 169, 60)     # 法老琥珀金 #E5A93C
COLOR_NEUTRAL = RGBColor(31, 36, 33)        # 曜石深灰 #1F2421
COLOR_MUTED = RGBColor(107, 114, 108)       # 沉香灰 #6B726C
COLOR_SURFACE = RGBColor(255, 255, 255)     # 純白 #FFFFFF
COLOR_SURFACE_WARM = RGBColor(250, 247, 242)# 羊皮紙暖白 #FAF7F2
COLOR_BORDER = RGBColor(232, 223, 216)      # 暖沙灰 #E8DFD8
COLOR_ACCENT_SOFT = RGBColor(254, 247, 233) # 琥珀柔光 #FEF7E9

FONT_HEADING = "Outfit"
FONT_BODY = "Inter"

OUTPUT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_FILE = os.path.join(OUTPUT_DIR, "Cloud_Architecture_Strategic_Deck_HowardLiao.pptx")

def create_cloud_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    def add_footer(slide, cur, total=4):
        fline = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.9), Inches(11.733), Inches(0.02))
        fline.fill.solid()
        fline.fill.fore_color.rgb = COLOR_BORDER
        fline.line.color.rgb = COLOR_BORDER

        lbox = slide.shapes.add_textbox(Inches(0.8), Inches(6.95), Inches(7.0), Inches(0.4))
        tf_l = lbox.text_frame
        p_l = tf_l.paragraphs[0]
        p_l.text = "報告人 Howard Liao Ph.D.(廖倫豪 博士) | 企業級混合多雲架構戰略"
        p_l.font.size = Pt(10)
        p_l.font.color.rgb = COLOR_MUTED
        p_l.font.name = FONT_BODY

        rbox = slide.shapes.add_textbox(Inches(10.5), Inches(6.95), Inches(2.0), Inches(0.4))
        tf_r = rbox.text_frame
        p_r = tf_r.paragraphs[0]
        p_r.text = f"{cur} / {total}"
        p_r.alignment = PP_ALIGN.RIGHT
        p_r.font.size = Pt(10)
        p_r.font.color.rgb = COLOR_MUTED
        p_r.font.name = FONT_BODY

    # ==========================================
    # SLIDE 1: 封面 (Cover)
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)

    # Accent bar
    bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.8), Inches(0.2), Inches(5.9))
    bar.fill.solid()
    bar.fill.fore_color.rgb = COLOR_PRIMARY
    bar.line.fill.background()

    gold_bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.1), Inches(0.8), Inches(0.08), Inches(2.2))
    gold_bar.fill.solid()
    gold_bar.fill.fore_color.rgb = COLOR_TERTIARY
    gold_bar.line.fill.background()

    tb1 = s1.shapes.add_textbox(Inches(1.5), Inches(1.6), Inches(10.5), Inches(3.2))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p_tag = tf1.paragraphs[0]
    p_tag.text = "ENTERPRISE CLOUD STRATEGY & GOVERNANCE"
    p_tag.font.size = Pt(12)
    p_tag.font.bold = True
    p_tag.font.color.rgb = COLOR_SECONDARY
    p_tag.font.name = FONT_HEADING
    p_tag.space_after = Pt(14)

    p_title = tf1.add_paragraph()
    p_title.text = "新世代企業混合多雲架構規劃"
    p_title.font.size = Pt(38)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_PRIMARY
    p_title.font.name = FONT_HEADING
    p_title.space_after = Pt(12)

    p_sub = tf1.add_paragraph()
    p_sub.text = "高可用性、彈性擴展與 FinOps 成本最佳化的現代化雲端落地藍圖"
    p_sub.font.size = Pt(18)
    p_sub.font.color.rgb = COLOR_NEUTRAL
    p_sub.font.name = FONT_BODY

    # Author Card
    auth = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(5.2), Inches(6.8), Inches(1.15))
    auth.fill.solid()
    auth.fill.fore_color.rgb = COLOR_ACCENT_SOFT
    auth.line.color.rgb = COLOR_TERTIARY
    auth.line.width = Pt(1.5)

    tf_a = auth.text_frame
    tf_a.margin_left = Inches(0.3)
    tf_a.margin_top = Inches(0.2)
    pa1 = tf_a.paragraphs[0]
    pa1.text = "報告人 Howard Liao Ph.D.(廖倫豪 博士)"
    pa1.font.size = Pt(15)
    pa1.font.bold = True
    pa1.font.color.rgb = COLOR_PRIMARY
    pa1.font.name = FONT_HEADING

    pa2 = tf_a.add_paragraph()
    pa2.text = "資深雲端架構師 / 資訊工程博士  |  2026 戰略規劃專案"
    pa2.font.size = Pt(11)
    pa2.font.color.rgb = COLOR_SECONDARY
    pa2.font.name = FONT_BODY

    # ==========================================
    # SLIDE 2: 雲端架構五大核心支柱 (Well-Architected Framework)
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    add_footer(s2, 2)

    hb2 = s2.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.7), Inches(1.0))
    p2_tag = hb2.text_frame.paragraphs[0]
    p2_tag.text = "WELL-ARCHITECTED FRAMEWORK"
    p2_tag.font.size = Pt(11)
    p2_tag.font.bold = True
    p2_tag.font.color.rgb = COLOR_SECONDARY
    p2_tag.font.name = FONT_HEADING

    p2_title = hb2.text_frame.add_paragraph()
    p2_title.text = "雲端架構五大支柱與治理矩陣"
    p2_title.font.size = Pt(26)
    p2_title.font.bold = True
    p2_title.font.color.rgb = COLOR_PRIMARY
    p2_title.font.name = FONT_HEADING

    pillars = [
        ("01", "卓越營運 (Ops)", "基礎架構即程式碼 (IaC)，端到端 CI/CD 全自動部署與微服務觀測。"),
        ("02", "安全性 (Security)", "零信任身份鑑別、密鑰統一金庫 (Vault) 管理與靜態/傳輸全面加密。"),
        ("03", "高可用可靠性 (Resilience)", "跨區域多活容災架構 (Active-Active)、自動水平彈性伸縮 (HPA)。"),
        ("04", "效能效率 (Efficiency)", "無伺服器架構 (Serverless)、全球邊緣 CDN 加速與微服務容器化。"),
        ("05", "FinOps 成本優化", "精確標籤 (Tagging)、雲端保留容量採購與動態負載調度減省 35%+。")
    ]

    card_w = Inches(2.18)
    card_h = Inches(4.3)
    start_x = Inches(0.8)
    gap = Inches(0.2)
    y_pos = Inches(2.1)

    for i, (num, title, desc) in enumerate(pillars):
        cx = start_x + i * (card_w + gap)
        card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, y_pos, card_w, card_h)
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_SURFACE_WARM
        card.line.color.rgb = COLOR_BORDER
        card.line.width = Pt(1)

        tfc = card.text_frame
        tfc.margin_left = Inches(0.18)
        tfc.margin_right = Inches(0.18)
        tfc.margin_top = Inches(0.25)
        tfc.word_wrap = True

        pn = tfc.paragraphs[0]
        pn.text = num
        pn.font.size = Pt(24)
        pn.font.bold = True
        pn.font.color.rgb = COLOR_TERTIARY
        pn.font.name = FONT_HEADING
        pn.space_after = Pt(8)

        pt = tfc.add_paragraph()
        pt.text = title
        pt.font.size = Pt(14)
        pt.font.bold = True
        pt.font.color.rgb = COLOR_SECONDARY
        pt.font.name = FONT_HEADING
        pt.space_after = Pt(12)

        pd = tfc.add_paragraph()
        pd.text = desc
        pd.font.size = Pt(11)
        pd.font.color.rgb = COLOR_NEUTRAL
        pd.font.name = FONT_BODY

    # ==========================================
    # SLIDE 3: 混合多雲拓撲與落地規劃 (Topology & Comparison)
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    add_footer(s3, 3)

    hb3 = s3.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.7), Inches(1.0))
    p3_tag = hb3.text_frame.paragraphs[0]
    p3_tag.text = "MODERN CLOUD TOPOLOGY"
    p3_tag.font.size = Pt(11)
    p3_tag.font.bold = True
    p3_tag.font.color.rgb = COLOR_SECONDARY

    p3_title = hb3.text_frame.add_paragraph()
    p3_title.text = "地端核心與多雲彈性拓撲分工"
    p3_title.font.size = Pt(26)
    p3_title.font.bold = True
    p3_title.font.color.rgb = COLOR_PRIMARY

    w_col = Inches(5.7)
    h_col = Inches(4.3)

    # Left: Private Cloud
    c_left = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.1), w_col, h_col)
    c_left.fill.solid()
    c_left.fill.fore_color.rgb = COLOR_SURFACE_WARM
    c_left.line.color.rgb = COLOR_BORDER

    tfl = c_left.text_frame
    tfl.margin_left = Inches(0.3)
    tfl.margin_top = Inches(0.3)
    tfl.word_wrap = True

    plt = tfl.paragraphs[0]
    plt.text = "地端私有核心 (On-Premises / Private)"
    plt.font.size = Pt(18)
    plt.font.bold = True
    plt.font.color.rgb = COLOR_PRIMARY
    plt.space_after = Pt(12)

    l_items = [
        "核心機敏資料庫與法規監管資料存儲（金融/個資合規隔離）。",
        "超低延遲專線 Direct Connect / ExpressRoute 互連。",
        "自建 Kubernetes 叢集 (OpenShift/Rancher) 運行關鍵核心交易。",
        "硬體安全模組 (HSM) 實體金鑰託管與私有金鑰隔離。"
    ]
    for item in l_items:
        p = tfl.add_paragraph()
        p.text = f"•  {item}"
        p.font.size = Pt(13)
        p.font.color.rgb = COLOR_NEUTRAL
        p.font.name = FONT_BODY
        p.space_after = Pt(8)

    # Right: Public Cloud
    c_right = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(2.1), w_col, h_col)
    c_right.fill.solid()
    c_right.fill.fore_color.rgb = COLOR_ACCENT_SOFT
    c_right.line.color.rgb = COLOR_TERTIARY
    c_right.line.width = Pt(2)

    tfr = c_right.text_frame
    tfr.margin_left = Inches(0.3)
    tfr.margin_top = Inches(0.3)
    tfr.word_wrap = True

    prt = tfr.paragraphs[0]
    prt.text = "公有雲彈性層 (Multi-Cloud Elastic Tier)"
    prt.font.size = Pt(18)
    prt.font.bold = True
    prt.font.color.rgb = COLOR_SECONDARY
    prt.space_after = Pt(12)

    r_items = [
        "面對千萬用戶的邊緣端點、動態網站與 API Gateway 集群。",
        "AI/ML 大規模分散式訓練與大數據批次計算彈性節點。",
        "跨區域跨可用區 (Multi-AZ) 自動故障轉移與冷熱資料分層存儲。",
        "FinOps 動態 Spot / Reserved Instance 智能調度降低支出。"
    ]
    for item in r_items:
        p = tfr.add_paragraph()
        p.text = f"✔  {item}"
        p.font.size = Pt(13)
        p.font.color.rgb = COLOR_NEUTRAL
        p.font.name = FONT_BODY
        p.space_after = Pt(8)

    # ==========================================
    # SLIDE 4: 戰略推進與總結 (Roadmap & Sign-off)
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    add_footer(s4, 4)

    center_card = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(1.4), Inches(10.333), Inches(5.0))
    center_card.fill.solid()
    center_card.fill.fore_color.rgb = COLOR_SURFACE_WARM
    center_card.line.color.rgb = COLOR_TERTIARY
    center_card.line.width = Pt(1.5)

    tfc4 = center_card.text_frame
    tfc4.margin_top = Inches(0.5)
    tfc4.margin_left = Inches(0.6)
    tfc4.margin_right = Inches(0.6)
    tfc4.word_wrap = True

    pc1 = tfc4.paragraphs[0]
    pc1.text = "IMPLEMENTATION ROADMAP & COMMITMENT"
    pc1.font.size = Pt(12)
    pc1.font.bold = True
    pc1.alignment = PP_ALIGN.CENTER
    pc1.font.color.rgb = COLOR_SECONDARY

    pc2 = tfc4.add_paragraph()
    pc2.text = "三階段雲端轉型推動路徑"
    pc2.font.size = Pt(28)
    pc2.font.bold = True
    pc2.alignment = PP_ALIGN.CENTER
    pc2.font.color.rgb = COLOR_PRIMARY
    pc2.space_after = Pt(18)

    pc3 = tfc4.add_paragraph()
    pc3.text = "【第 1 階段：地基加固】雲端登陸區 (Landing Zone) 規格固化與網路混合專線建立\n【第 2 階段：容器現代化】應用微服務解耦、CI/CD 自動化與 FinOps 成本可視化\n【第 3 階段：智慧雲原生】AI Agent 深度融合、多雲動態負載平衡與自癒容災"
    pc3.font.size = Pt(14)
    pc3.alignment = PP_ALIGN.CENTER
    pc3.font.color.rgb = COLOR_NEUTRAL
    pc3.space_after = Pt(26)

    pc4 = tfc4.add_paragraph()
    pc4.text = "報告人 Howard Liao Ph.D.(廖倫豪 博士)"
    pc4.font.size = Pt(18)
    pc4.font.bold = True
    pc4.alignment = PP_ALIGN.CENTER
    pc4.font.color.rgb = COLOR_PRIMARY

    prs.save(OUTPUT_FILE)
    print(f"Cloud Architecture Deck generated successfully: {OUTPUT_FILE}")

if __name__ == "__main__":
    create_cloud_deck()
