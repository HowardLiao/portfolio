#!/usr/bin/env python3
"""
Generate Standard PowerPoint Deck adhering to Egyptian Pharaoh Academic Design System
Author: Howard Liao Ph.D. (廖倫豪 博士)
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# Design Tokens
COLOR_PRIMARY = RGBColor(217, 107, 39)       # 暖赭石橙 #D96B27
COLOR_SECONDARY = RGBColor(43, 95, 107)     # 綠松石青 #2B5F6B
COLOR_TERTIARY = RGBColor(229, 169, 60)     # 法老琥珀金 #E5A93C
COLOR_NEUTRAL = RGBColor(31, 36, 33)        # 曜石深灰 #1F2421
COLOR_MUTED = RGBColor(107, 114, 108)       # 沉香灰 #6B726C
COLOR_SURFACE = RGBColor(255, 255, 255)     # 純白 #FFFFFF
COLOR_SURFACE_WARM = RGBColor(250, 247, 242)# 羊皮紙暖白 #FAF7F2
COLOR_BORDER = RGBColor(232, 223, 216)      # 暖沙灰 #E8DFD8
COLOR_ACCENT_SOFT = RGBColor(254, 247, 233) # 琥珀柔光 #FEF7E9

FONT_HEADING = "Outfit"
FONT_BODY = "Inter"

OUTPUT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_FILE = os.path.join(OUTPUT_DIR, "Pharaoh_Academic_Presentation_HowardLiao.pptx")

def create_deck():
    prs = Presentation()
    # 16:9 Widescreen
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    def add_standard_footer(slide, current_page, total_pages=4):
        # Top subtle line
        footer_line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.9), Inches(11.733), Inches(0.02))
        footer_line.fill.solid()
        footer_line.fill.fore_color.rgb = COLOR_BORDER
        footer_line.line.color.rgb = COLOR_BORDER

        # Left attribution
        left_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.95), Inches(6.0), Inches(0.4))
        tf_l = left_box.text_frame
        tf_l.word_wrap = True
        p_l = tf_l.paragraphs[0]
        p_l.text = "報告人 Howard Liao Ph.D.(廖倫豪 博士) | 專業戰略與學術研究"
        p_l.font.size = Pt(10)
        p_l.font.color.rgb = COLOR_MUTED
        p_l.font.name = FONT_BODY

        # Right page number
        right_box = slide.shapes.add_textbox(Inches(10.5), Inches(6.95), Inches(2.0), Inches(0.4))
        tf_r = right_box.text_frame
        p_r = tf_r.paragraphs[0]
        p_r.text = f"{current_page} / {total_pages}"
        p_r.alignment = PP_ALIGN.RIGHT
        p_r.font.size = Pt(10)
        p_r.font.color.rgb = COLOR_MUTED
        p_r.font.name = FONT_BODY

    # ==========================================
    # SLIDE 1: 封面 (Title Slide)
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)

    # Left accent bar (Egyptian Pharaoh Ochre)
    accent_bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.8), Inches(0.2), Inches(5.9))
    accent_bar.fill.solid()
    accent_bar.fill.fore_color.rgb = COLOR_PRIMARY
    accent_bar.line.fill.background()

    # Small gold accent line
    gold_bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.1), Inches(0.8), Inches(0.08), Inches(2.0))
    gold_bar.fill.solid()
    gold_bar.fill.fore_color.rgb = COLOR_TERTIARY
    gold_bar.line.fill.background()

    # Main Title Box
    title_box = s1.shapes.add_textbox(Inches(1.5), Inches(1.8), Inches(10.5), Inches(3.0))
    tf1 = title_box.text_frame
    tf1.word_wrap = True

    # Category Tag
    p_tag = tf1.paragraphs[0]
    p_tag.text = "STRATEGIC RESEARCH & ARCHITECTURE DESIGN"
    p_tag.font.size = Pt(12)
    p_tag.font.bold = True
    p_tag.font.color.rgb = COLOR_SECONDARY
    p_tag.font.name = FONT_HEADING
    p_tag.space_after = Pt(14)

    # Big Title
    p_main = tf1.add_paragraph()
    p_main.text = "埃及法老溫暖學術風格：全鏈路設計系統"
    p_main.font.size = Pt(38)
    p_main.font.bold = True
    p_main.font.color.rgb = COLOR_PRIMARY
    p_main.font.name = FONT_HEADING
    p_main.space_after = Pt(12)

    # Subtitle
    p_sub = tf1.add_paragraph()
    p_sub.text = "跨 PPT、Word、Google Docs 及 Web 前端的一致性視覺規範實踐架構"
    p_sub.font.size = Pt(18)
    p_sub.font.color.rgb = COLOR_NEUTRAL
    p_sub.font.name = FONT_BODY

    # Attribution Box (Bottom Left - Official Format)
    author_card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(5.2), Inches(6.5), Inches(1.1))
    author_card.fill.solid()
    author_card.fill.fore_color.rgb = COLOR_ACCENT_SOFT
    author_card.line.color.rgb = COLOR_TERTIARY
    author_card.line.width = Pt(1.5)

    tf_auth = author_card.text_frame
    tf_auth.margin_left = Inches(0.3)
    tf_auth.margin_top = Inches(0.18)
    p_a1 = tf_auth.paragraphs[0]
    p_a1.text = "報告人 Howard Liao Ph.D.(廖倫豪 博士)"
    p_a1.font.size = Pt(15)
    p_a1.font.bold = True
    p_a1.font.color.rgb = COLOR_PRIMARY
    p_a1.font.name = FONT_HEADING

    p_a2 = tf_auth.add_paragraph()
    p_a2.text = "資深架構師 / 資訊安全長 / 數位轉型顧問"
    p_a2.font.size = Pt(11)
    p_a2.font.color.rgb = COLOR_SECONDARY
    p_a2.font.name = FONT_BODY

    # ==========================================
    # SLIDE 2: 目錄與核心框架 (Agenda / Pillars)
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    add_standard_footer(s2, 2)

    # Section header
    h2_box = s2.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.7), Inches(1.0))
    tf2_h = h2_box.text_frame
    p_sh = tf2_h.paragraphs[0]
    p_sh.text = "DESIGN ARCHITECTURE"
    p_sh.font.size = Pt(11)
    p_sh.font.bold = True
    p_sh.font.color.rgb = COLOR_SECONDARY
    p_sh.font.name = FONT_HEADING

    p_mt = tf2_h.add_paragraph()
    p_mt.text = "核心架構：四大一致性落地模組"
    p_mt.font.size = Pt(26)
    p_mt.font.bold = True
    p_mt.font.color.rgb = COLOR_PRIMARY
    p_mt.font.name = FONT_HEADING

    # 4 Pillar Cards
    pillars = [
        ("01", "Design Tokens", "定義暖赭橙、綠松青、琥珀金與曜石灰階，統一 JSON 與 CSS 格式。"),
        ("02", "Office Automation", "利用 python-pptx 與 python-docx 自動映射母片樣式與規格。"),
        ("03", "Web & Components", "提供卡片、微細邊框、表格與柔和陰影的專屬 CSS/Tailwind 類別。"),
        ("04", "Identity & Sign-off", "統一封面、頁尾與標題階層，固化博士署名與高階學術權威感。")
    ]

    card_w = Inches(2.7)
    card_h = Inches(4.2)
    start_x = Inches(0.8)
    gap = Inches(0.3)
    y_pos = Inches(2.1)

    for i, (num, title, desc) in enumerate(pillars):
        cx = start_x + i * (card_w + gap)
        card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, y_pos, card_w, card_h)
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_SURFACE_WARM
        card.line.color.rgb = COLOR_BORDER
        card.line.width = Pt(1)

        tf_c = card.text_frame
        tf_c.margin_left = Inches(0.25)
        tf_c.margin_right = Inches(0.25)
        tf_c.margin_top = Inches(0.3)
        tf_c.word_wrap = True

        # Number in Amber Gold
        p_n = tf_c.paragraphs[0]
        p_n.text = num
        p_n.font.size = Pt(28)
        p_n.font.bold = True
        p_n.font.color.rgb = COLOR_TERTIARY
        p_n.font.name = FONT_HEADING
        p_n.space_after = Pt(10)

        # Title in Turquoise Cyan
        p_t = tf_c.add_paragraph()
        p_t.text = title
        p_t.font.size = Pt(16)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_SECONDARY
        p_t.font.name = FONT_HEADING
        p_t.space_after = Pt(14)

        # Desc in Obsidian Dark Gray
        p_d = tf_c.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(12)
        p_d.font.color.rgb = COLOR_NEUTRAL
        p_d.font.name = FONT_BODY

    # ==========================================
    # SLIDE 3: 深度內容與卡片佈局 (Content Split)
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    add_standard_footer(s3, 3)

    # Header
    h3_box = s3.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.7), Inches(1.0))
    tf3_h = h3_box.text_frame
    p3_sh = tf3_h.paragraphs[0]
    p3_sh.text = "CROSS-PLATFORM STANDARDIZATION"
    p3_sh.font.size = Pt(11)
    p3_sh.font.bold = True
    p3_sh.font.color.rgb = COLOR_SECONDARY
    p3_sh.font.name = FONT_HEADING

    p3_mt = tf3_h.add_paragraph()
    p3_mt.text = "跨格式精確映射標準"
    p3_mt.font.size = Pt(26)
    p3_mt.font.bold = True
    p3_mt.font.color.rgb = COLOR_PRIMARY
    p3_mt.font.name = FONT_HEADING

    # Left Column: Standard Card
    left_w = Inches(5.7)
    left_h = Inches(4.3)
    c_left = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.1), left_w, left_h)
    c_left.fill.solid()
    c_left.fill.fore_color.rgb = COLOR_SURFACE_WARM
    c_left.line.color.rgb = COLOR_BORDER

    tf_l = c_left.text_frame
    tf_l.margin_left = Inches(0.3)
    tf_l.margin_top = Inches(0.3)
    tf_l.word_wrap = True

    p_lt = tf_l.paragraphs[0]
    p_lt.text = "規範要點 (System Rules)"
    p_lt.font.size = Pt(18)
    p_lt.font.bold = True
    p_lt.font.color.rgb = COLOR_PRIMARY
    p_lt.space_after = Pt(12)

    bullets = [
        "背景採用高雅純白與羊皮紙暖白，護眼且兼具極致列印友好性。",
        "內文改用曜石深灰 (#1F2421)，徹底告別生硬刺眼的純黑字體。",
        "圖表與標籤採用尼羅河綠松石青 (#2B5F6B)，塑造冷靜權威感。",
        "所有重要決策點與核心 KPI 一律套用琥珀金 (#E5A93C) 亮點框。"
    ]
    for b in bullets:
        pb = tf_l.add_paragraph()
        pb.text = f"•  {b}"
        pb.font.size = Pt(13)
        pb.font.color.rgb = COLOR_NEUTRAL
        pb.font.name = FONT_BODY
        pb.space_after = Pt(8)

    # Right Column: Highlight Card
    c_right = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(2.1), left_w, left_h)
    c_right.fill.solid()
    c_right.fill.fore_color.rgb = COLOR_ACCENT_SOFT
    c_right.line.color.rgb = COLOR_TERTIARY
    c_right.line.width = Pt(2)

    tf_r = c_right.text_frame
    tf_r.margin_left = Inches(0.3)
    tf_r.margin_top = Inches(0.3)
    tf_r.word_wrap = True

    p_rt = tf_r.paragraphs[0]
    p_rt.text = "關鍵效益 (Strategic Value)"
    p_rt.font.size = Pt(18)
    p_rt.font.bold = True
    p_rt.font.color.rgb = COLOR_SECONDARY
    p_rt.space_after = Pt(12)

    benefits = [
        "100% 品牌語彙一致：無論是產出 PPT、Word 還是網頁，閱讀體驗全然一體。",
        "單一真實來源 (Single Source)：修改 DESIGN.md 即可透過管線連動更新所有模板。",
        "學術與顧問公信力：暖系法老色調帶來莊重、高品味與清晰的論證節奏。",
        "即時可自動化：由 HermesAgent 隨時調用 Python 腳本產生原生標準檔。"
    ]
    for b in benefits:
        pb = tf_r.add_paragraph()
        pb.text = f"✔  {b}"
        pb.font.size = Pt(13)
        pb.font.color.rgb = COLOR_NEUTRAL
        pb.font.name = FONT_BODY
        pb.space_after = Pt(8)

    # ==========================================
    # SLIDE 4: 總結與專屬署名 (Conclusion)
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    add_standard_footer(s4, 4)

    center_card = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.8), Inches(1.5), Inches(9.733), Inches(4.8))
    center_card.fill.solid()
    center_card.fill.fore_color.rgb = COLOR_SURFACE_WARM
    center_card.line.color.rgb = COLOR_TERTIARY
    center_card.line.width = Pt(1.5)

    tf_cen = center_card.text_frame
    tf_cen.margin_top = Inches(0.6)
    tf_cen.margin_left = Inches(0.5)
    tf_cen.margin_right = Inches(0.5)
    tf_cen.word_wrap = True

    p_c1 = tf_cen.paragraphs[0]
    p_c1.text = "CONCLUSION & NEXT STEPS"
    p_c1.font.size = Pt(12)
    p_c1.font.bold = True
    p_c1.alignment = PP_ALIGN.CENTER
    p_c1.font.color.rgb = COLOR_SECONDARY

    p_c2 = tf_cen.add_paragraph()
    p_c2.text = "全面啟動統一視覺規範"
    p_c2.font.size = Pt(30)
    p_c2.font.bold = True
    p_c2.alignment = PP_ALIGN.CENTER
    p_c2.font.color.rgb = COLOR_PRIMARY
    p_c2.space_after = Pt(16)

    p_c3 = tf_cen.add_paragraph()
    p_c3.text = "後續在 HermesAgent 生成的所有研究簡報、Word 分析專案、互動網頁與架構圖，\n皆自動遵從本埃及法老溫暖學術規範，實現最高品質的專業品牌一致性。"
    p_c3.font.size = Pt(15)
    p_c3.alignment = PP_ALIGN.CENTER
    p_c3.font.color.rgb = COLOR_NEUTRAL
    p_c3.space_after = Pt(30)

    # Formal Signature inside Box
    p_c4 = tf_cen.add_paragraph()
    p_c4.text = "報告人 Howard Liao Ph.D.(廖倫豪 博士)"
    p_c4.font.size = Pt(18)
    p_c4.font.bold = True
    p_c4.alignment = PP_ALIGN.CENTER
    p_c4.font.color.rgb = COLOR_PRIMARY

    prs.save(OUTPUT_FILE)
    print(f"PPTX generated successfully at: {OUTPUT_FILE}")

if __name__ == "__main__":
    create_deck()
